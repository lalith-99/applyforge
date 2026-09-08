#!/usr/bin/env python3
"""Aggregate official DOL LCA/PERM disclosure workbooks and import evidence.

This tool intentionally lives outside the production API/AI-worker runtime.
It reads locally downloaded OFLC XLSX disclosure files, aggregates employer-level
historical evidence, and sends small batches to ApplyForge's protected admin endpoint.

DOL disclosure data is historical evidence only. It must never be interpreted
as a guarantee that a specific current role provides immigration sponsorship.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import urllib.error
import urllib.request
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


EMPLOYER_HEADERS = (
    # LCA / older disclosure schemas
    "EMPLOYER_NAME",
    "EMPLOYER_BUSINESS_NAME",

    # Current PERM ETA-9089 schema
    "EMP_BUSINESS_NAME",
)
STATUS_HEADERS = ("CASE_STATUS", "STATUS")
VISA_HEADERS = ("VISA_CLASS", "VISA_TYPE")
DECISION_HEADERS = (
    "DECISION_DATE",
    "CASE_DECISION_DATE",
    "DETERMINATION_DATE",
)

LEGAL_SUFFIXES = {
    "inc",
    "incorporated",
    "llc",
    "corp",
    "corporation",
    "ltd",
    "limited",
    "pllc",
    "lp",
    "llp",
    "co",
    "company",
}


@dataclass
class Aggregate:
    employer_name: str
    certified_count: int = 0
    denied_count: int = 0
    withdrawn_count: int = 0
    other_count: int = 0
    total_count: int = 0
    latest_decision_date: date | None = None


def normalize_employer(name: str) -> str:
    words = re.sub(r"[^a-z0-9]+", " ", name.strip().lower()).split()

    while len(words) > 1 and words[-1] in LEGAL_SUFFIXES:
        words.pop()

    return " ".join(words)


def read_file(path: str) -> Path:
    """Validate and return a local DOL disclosure workbook path."""

    file_path = Path(path).expanduser().resolve()

    if not file_path.exists():
        raise FileNotFoundError(
            f"DOL disclosure file not found: {file_path}"
        )

    if not file_path.is_file():
        raise ValueError(
            f"DOL disclosure path is not a file: {file_path}"
        )

    if file_path.suffix.lower() != ".xlsx":
        raise ValueError(
            f"DOL disclosure file must be an .xlsx workbook: {file_path}"
        )

    return file_path


def header_index(
    headers: list[str],
    candidates: Iterable[str],
    required: bool = True,
) -> int | None:
    normalized = {
        str(value).strip().upper(): i
        for i, value in enumerate(headers)
    }

    for candidate in candidates:
        if candidate in normalized:
            return normalized[candidate]

    if required:
        raise ValueError(
            f"missing required header; expected one of {tuple(candidates)}"
        )

    return None


def parse_date(value: Any) -> date | None:
    if value is None or value == "":
        return None

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    text = str(value).strip()

    for fmt in (
        "%Y-%m-%d",
        "%m/%d/%Y",
        "%m/%d/%y",
        "%Y/%m/%d",
    ):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass

    return None


def is_h1b(value: Any) -> bool:
    if value is None:
        return False

    normalized = re.sub(
        r"[^A-Z0-9]+",
        "",
        str(value).upper(),
    )

    return normalized == "H1B"


def classify_status(status: str) -> tuple[bool, bool, bool, bool]:
    value = status.strip().upper()

    certified = "CERTIFIED" in value
    denied = "DENIED" in value
    withdrawn = "WITHDRAWN" in value
    other = not (certified or denied or withdrawn)

    return certified, denied, withdrawn, other


def aggregate_workbook(
    file_path: Path,
    program: str,
) -> dict[str, Aggregate]:
    print(
        f"{program}: opening workbook {file_path}",
        file=sys.stderr,
    )

    workbook = load_workbook(
        filename=file_path,
        read_only=True,
        data_only=True,
    )

    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)

        try:
            raw_headers = next(rows)
        except StopIteration as exc:
            raise ValueError(
                f"workbook has no rows: {file_path}"
            ) from exc

        headers = [
            str(value).strip() if value is not None else ""
            for value in raw_headers
        ]

        employer_idx = header_index(
            headers,
            EMPLOYER_HEADERS,
        )

        status_idx = header_index(
            headers,
            STATUS_HEADERS,
        )

        decision_idx = header_index(
            headers,
            DECISION_HEADERS,
            required=False,
        )

        visa_idx = header_index(
            headers,
            VISA_HEADERS,
            required=False,
        )

        aggregate: dict[str, Aggregate] = {}

        for row in rows:
            employer = str(
                row[employer_idx] or ""
            ).strip()

            if not employer:
                continue

            if (
                program == "LCA_H1B"
                and visa_idx is not None
                and not is_h1b(row[visa_idx])
            ):
                continue

            status = str(
                row[status_idx] or ""
            ).strip()

            key = normalize_employer(employer)

            if not key:
                continue

            item = aggregate.get(key)

            if item is None:
                item = Aggregate(
                    employer_name=employer
                )
                aggregate[key] = item

            item.total_count += 1

            certified, denied, withdrawn, other = (
                classify_status(status)
            )

            # These are evidence dimensions rather than mutually
            # exclusive buckets.
            #
            # For example, "Certified-Withdrawn" records both facts.
            item.certified_count += int(certified)
            item.denied_count += int(denied)
            item.withdrawn_count += int(withdrawn)
            item.other_count += int(other)

            if decision_idx is not None:
                decision = parse_date(
                    row[decision_idx]
                )

                if (
                    decision
                    and (
                        item.latest_decision_date is None
                        or decision
                        > item.latest_decision_date
                    )
                ):
                    item.latest_decision_date = decision

        return aggregate

    finally:
        workbook.close()


def make_rows(
    aggregates: dict[str, Aggregate],
    program: str,
    fiscal_year: int,
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []

    for item in aggregates.values():
        rows.append(
            {
                "employer_name": item.employer_name,
                "program": program,
                "fiscal_year": fiscal_year,
                "certified_count": item.certified_count,
                "denied_count": item.denied_count,
                "withdrawn_count": item.withdrawn_count,
                "other_count": item.other_count,
                "total_count": item.total_count,
                "latest_decision_date": (
                    f"{item.latest_decision_date.isoformat()}T00:00:00Z"
                    if item.latest_decision_date
                    else None
                ),
            }
        )

    rows.sort(
        key=lambda item: (
            -item["certified_count"],
            item["employer_name"].lower(),
        )
    )

    return rows


def post_batch(
    api_base: str,
    admin_token: str,
    payload: dict[str, Any],
) -> dict[str, Any]:
    url = (
        api_base.rstrip("/")
        + "/admin/immigration/evidence/import"
    )

    body = json.dumps(payload).encode("utf-8")

    request = urllib.request.Request(
        url,
        data=body,
        method="POST",
        headers={
            "Content-Type": "application/json",
            "X-ApplyForge-Admin-Token": admin_token,
            "User-Agent": "ApplyForge-DOL-Importer/1.0",
        },
    )

    try:
        with urllib.request.urlopen(
            request,
            timeout=120,
        ) as response:
            return json.loads(
                response.read().decode("utf-8")
            )

    except urllib.error.HTTPError as exc:
        details = exc.read().decode(
            "utf-8",
            errors="replace",
        )

        raise RuntimeError(
            f"ApplyForge import failed "
            f"with HTTP {exc.code}: {details}"
        ) from exc

    except urllib.error.URLError as exc:
        raise RuntimeError(
            f"Could not connect to ApplyForge API "
            f"at {url}: {exc.reason}"
        ) from exc


def chunks(
    rows: list[dict[str, Any]],
    size: int,
) -> Iterable[list[dict[str, Any]]]:
    for start in range(0, len(rows), size):
        yield rows[start:start + size]


def import_program(
    *,
    program: str,
    source_file: str,
    fiscal_year: int,
    source_release: str,
    api_base: str,
    admin_token: str,
    batch_size: int,
    dry_run: bool,
) -> tuple[int, int]:

    file_path = read_file(source_file)

    print(
        f"Reading {program} disclosure: {file_path}",
        file=sys.stderr,
    )

    aggregates = aggregate_workbook(
        file_path,
        program,
    )

    rows = make_rows(
        aggregates,
        program,
        fiscal_year,
    )

    certified = sum(
        row["certified_count"]
        for row in rows
    )

    print(
        f"{program}: "
        f"{len(rows):,} employers, "
        f"{certified:,} certified-case signals",
        file=sys.stderr,
    )

    if dry_run:
        return len(rows), certified

    imported = 0

    for batch_number, batch in enumerate(
        chunks(rows, batch_size),
        start=1,
    ):
        response = post_batch(
            api_base,
            admin_token,
            {
                "source_release": source_release,

                # Preserve provenance even though this is
                # now imported from a local file.
                "source_url": file_path.as_uri(),

                "rows": batch,
            },
        )

        imported += int(
            response.get("imported", 0)
        )

        print(
            f"{program}: "
            f"batch {batch_number} completed; "
            f"imported {imported:,}/{len(rows):,} employers",
            file=sys.stderr,
        )

    return imported, certified


def main() -> int:
    parser = argparse.ArgumentParser(
        description=(
            "Import locally downloaded DOL "
            "LCA/PERM disclosure workbooks into ApplyForge."
        )
    )

    parser.add_argument(
        "--lca-file",
        default=str(
            Path.home()
            / "Downloads"
            / "LCA_Disclosure_Data_FY2026_Q3.xlsx"
        ),
        help="Path to the DOL LCA disclosure XLSX file",
    )

    parser.add_argument(
        "--perm-file",
        default=str(
            Path.home()
            / "Downloads"
            / "PERM_Disclosure_Data_FY2026_Q3.xlsx"
        ),
        help="Path to the DOL PERM disclosure XLSX file",
    )

    parser.add_argument(
        "--fiscal-year",
        type=int,
        default=2026,
    )

    parser.add_argument(
        "--source-release",
        default="FY2026_Q3",
    )

    parser.add_argument(
        "--api-base",
        default="http://localhost:8080/api/v1",
    )

    parser.add_argument(
        "--admin-token",
        default="",
    )

    parser.add_argument(
        "--batch-size",
        type=int,
        default=1000,
    )

    parser.add_argument(
        "--dry-run",
        action="store_true",
    )

    args = parser.parse_args()

    if not args.lca_file and not args.perm_file:
        parser.error(
            "at least one of --lca-file "
            "or --perm-file is required"
        )

    if not args.dry_run and not args.admin_token:
        parser.error(
            "--admin-token is required "
            "unless --dry-run is used"
        )

    if args.batch_size < 1 or args.batch_size > 2000:
        parser.error(
            "--batch-size must be between 1 and 2000"
        )

    total_employers = 0
    total_certified = 0

    if args.lca_file:
        employers, certified = import_program(
            program="LCA_H1B",
            source_file=args.lca_file,
            fiscal_year=args.fiscal_year,
            source_release=args.source_release,
            api_base=args.api_base,
            admin_token=args.admin_token,
            batch_size=args.batch_size,
            dry_run=args.dry_run,
        )

        total_employers += employers
        total_certified += certified

    if args.perm_file:
        employers, certified = import_program(
            program="PERM",
            source_file=args.perm_file,
            fiscal_year=args.fiscal_year,
            source_release=args.source_release,
            api_base=args.api_base,
            admin_token=args.admin_token,
            batch_size=args.batch_size,
            dry_run=args.dry_run,
        )

        total_employers += employers
        total_certified += certified

    print(
        json.dumps(
            {
                "employer_program_rows": total_employers,
                "certified_case_signals": total_certified,
                "dry_run": args.dry_run,
            },
            indent=2,
        )
    )

    return 0


if __name__ == "__main__":
    raise SystemExit(main())